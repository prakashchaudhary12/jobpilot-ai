from docx import Document
from docx.shared import Pt


def create_docx_file(content, file_path):
    document = Document()

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    for line in content.split("\n"):
        line = line.strip()

        if not line:
            document.add_paragraph()
            continue

        if line.startswith("#"):
            heading = line.replace("#", "").strip()
            document.add_heading(heading, level=1)

        elif line.startswith("##"):
            heading = line.replace("##", "").strip()
            document.add_heading(heading, level=2)

        elif line.startswith("-"):
            document.add_paragraph(
                line[1:].strip(),
                style="List Bullet",
            )

        else:
            document.add_paragraph(line)

    document.save(file_path)

    return file_path